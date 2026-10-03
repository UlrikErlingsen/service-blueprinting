"""Bounded spreadsheet reads, explicit mapping and portable Excel workbooks."""
from copy import deepcopy
import csv
from datetime import date, datetime
from io import BytesIO, StringIO
import math
from pathlib import Path
import re
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd

from . import input_format as fmt, limits, portable as io

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Reading has no limits on the user's own computer. In the public demo (SIGNAL_PUBLIC=1) the caps in limits.py stop
# reading early; each refusal says it is a demo limit.


def _too_many(count, cap):
    return cap is not None and count > cap


def _shape_problem(title, rows=None, columns=None):
    caps = []
    if rows is not None:
        caps.append(f"{limits.rows():,} rows")
    if columns is not None:
        caps.append(f"{limits.columns()} columns")
    return io.DataProblem(limits.demo(f"{title}: the demo reads at most {' and '.join(caps)} per sheet."))
COMMON_LABELS = {"id": "Reference", "name": "Name", "source_id": "Source reference", "note": "Notes", "url": "Source URL", "title": "Source title"}


def normalize(value):
    return re.sub(r"[^a-z0-9]+", "", str(value).casefold())


def label(field):
    return fmt.LABELS.get(field, COMMON_LABELS.get(field, field.replace("_", " ").capitalize()))


def suggest(columns, names):
    """Return only an unambiguous match; never guess the first column."""
    for name in names:
        matches = [c for c in columns if normalize(c) == normalize(name)]
        if len(matches) == 1:
            return matches[0]
    return None


def _frame(rows, title):
    rows = list(rows)
    while rows and all(v is None or v == "" for v in rows[-1]):
        rows.pop()
    if not rows:
        return None
    header = ["" if v is None else str(v).strip() for v in rows[0]]
    while header and header[-1] == "" and all(len(row) <= len(header)-1 or row[len(header)-1] in (None, "") for row in rows[1:]):
        header.pop()
    if not header or any(not c for c in header):
        raise io.DataProblem(f"{title}: put a name in every used column of the first row.")
    if len(set(map(normalize, header))) != len(header):
        raise io.DataProblem(f"{title}: column names must be distinct. Rename the duplicate headings.")
    if _too_many(len(rows) - 1, limits.rows()) or _too_many(len(header), limits.columns()):
        raise _shape_problem(title, rows=True, columns=True)
    data = []
    for index, row in enumerate(rows[1:], 2):
        if any(v not in (None, "") for v in row[len(header):]):
            raise io.DataProblem(f"{title}, row {index}: more values than column headings.")
        values = list(row[:len(header)]) + [None] * max(0, len(header)-len(row))
        if any(v not in (None, "") for v in values):
            data.append(values)
    return pd.DataFrame(data, columns=header, dtype=object)


def load_tables(files):
    """files is [(filename, bytes)]; no file contents enter a shared cache. Out of memory becomes a plain message."""
    try:
        return _load_tables(files)
    except MemoryError as exc:
        raise io.DataProblem(limits.MEMORY_NOTE) from exc


def _load_tables(files):
    if not files:
        raise io.DataProblem("Choose Excel (.xlsx) or UTF-8 CSV files.")
    cap = limits.upload_bytes()
    if _too_many(sum(len(raw) for _, raw in files), cap):
        raise io.DataProblem(limits.demo(f"Choose files totalling no more than {limits.megabytes(cap)}."))
    tables = {}
    cell_count = 0
    for filename, raw in files:
        suffix = Path(filename).suffix.lower()
        parsed = {}
        try:
            if suffix == ".csv":
                text = raw.decode("utf-8-sig")
                try:
                    dialect = csv.Sniffer().sniff(text[:16000], delimiters=",;\t")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.reader(StringIO(text), dialect, strict=True)
                rows = []
                for row in reader:
                    if _too_many(len(rows), limits.rows()) or _too_many(len(row), limits.columns()):
                        raise _shape_problem(filename, rows=True, columns=True)
                    rows.append(row)
                parsed[Path(filename).stem] = _frame(rows, filename)
            elif suffix == ".xlsx":
                with ZipFile(BytesIO(raw)) as archive:
                    if (_too_many(sum(f.file_size for f in archive.infolist()), limits.unzipped_bytes())
                            or _too_many(len(archive.infolist()), limits.zip_entries())):
                        raise io.DataProblem(limits.demo("This workbook is too large when opened for the demo."))
                formulas = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
                cached = None
                try:
                    if _too_many(len(formulas.worksheets), limits.sheets()):
                        raise io.DataProblem(limits.demo(f"The demo reads at most {limits.sheets()} sheets per workbook."))
                    cached = load_workbook(BytesIO(raw), read_only=True, data_only=True, keep_links=False)
                    for sheet in formulas.worksheets:
                        cap_rows = None if limits.rows() is None else limits.rows() + 1
                        if (sheet.max_row and _too_many(sheet.max_row, cap_rows)
                                or sheet.max_column and _too_many(sheet.max_column, limits.columns())):
                            raise _shape_problem(sheet.title, rows=True, columns=True)
                        rows = []
                        for row, saved in zip(sheet.iter_rows(), cached[sheet.title].iter_rows()):
                            if _too_many(len(rows), limits.rows()):
                                raise _shape_problem(sheet.title, rows=True)
                            values = []
                            for cell, value in zip(row, saved):
                                if cell.data_type == "f":
                                    if value.value is None:
                                        raise io.DataProblem(f"{sheet.title}, {cell.coordinate}: this formula has no saved result. Recalculate and save in Excel, or paste its values before uploading.")
                                    values.append(value.value)
                                elif cell.data_type == "e":
                                    raise io.DataProblem(f"{sheet.title}, {cell.coordinate}: fix the Excel error before importing.")
                                else:
                                    values.append(cell.value)
                            rows.append(values)
                        parsed[sheet.title] = _frame(rows, sheet.title)
                finally:
                    formulas.close()
                    if cached is not None:
                        cached.close()
            else:
                raise io.DataProblem("Use Excel .xlsx or CSV. For older .xls files, choose Save As → Excel Workbook (.xlsx).")
        except io.DataProblem:
            raise
        except (ValueError, TypeError, UnicodeError, csv.Error, BadZipFile, KeyError, OSError, SyntaxError) as exc:
            raise io.DataProblem(f"Could not read {filename}. Save a clean .xlsx or UTF-8 CSV with headings in the first row.") from exc
        for name, frame in parsed.items():
            if frame is None or normalize(name) in {"readme", "instructions"}:
                continue
            cell_count += (len(frame)+1) * len(frame.columns)
            if _too_many(cell_count, limits.cells()):
                raise io.DataProblem(limits.demo(f"The demo reads at most {limits.cells():,} cells across all files."))
            title = name if name not in tables else f"{Path(filename).stem} / {name}"
            if title in tables:
                raise io.DataProblem("Two files have the same sheet names. Rename a file or sheet to distinguish them.")
            tables[title] = frame
    if not tables:
        raise io.DataProblem("No data tables found. Put column headings on the first row and data underneath.")
    return tables


def number(value, comma=False, probability=False):
    if isinstance(value, bool):
        raise io.DataProblem("Use a number, not TRUE or FALSE.")
    if isinstance(value, str):
        value = value.strip()
        percent = value.endswith("%")
        if percent:
            if not probability:
                raise io.DataProblem("A percent sign is only supported for probability columns.")
            value = value[:-1].strip()
        if comma:
            if "." in value:
                raise io.DataProblem("This text uses a dot. Choose dot decimals or remove thousands separators.")
            value = value.replace(",", ".")
        elif "," in value:
            raise io.DataProblem("Choose comma decimals for values such as 1,5. Do not include thousands separators.")
        try:
            result = float(value) / (100 if percent else 1)
        except ValueError as exc:
            raise io.DataProblem("Use a numeric value; leave unknown values blank.") from exc
    else:
        try:
            result = float(value)
        except (ValueError, TypeError) as exc:
            raise io.DataProblem("Use a numeric value; leave unknown values blank.") from exc
    if not math.isfinite(result):
        raise io.DataProblem("Use a finite number; leave unknown values blank.")
    if probability and not 0 <= result <= 1:
        raise io.DataProblem("Probabilities use 0 to 1, or a percent sign (for example 40%).")
    return result


def missing(value):
    return value is None or isinstance(value, str) and not value.strip() or bool(pd.isna(value))


def cast(value, kind, comma=False):
    empty = missing(value)
    if kind == "basis" and empty:
        return "assumed"
    if kind.startswith("optional_") and empty:
        return "" if kind == "optional_text" else None
    if empty:
        raise io.DataProblem("Choose a column and fill in this value.")
    if kind in {"number", "optional_number", "optional_probability"}:
        return number(value, comma, "probability" in kind)
    text = str(value).strip()
    choices = {
        "lane": {"evidence": ["physical / digital evidence", "physical evidence", "digital evidence"], "customer": ["customer actions", "customer action"], "frontstage": ["visible staff / technology", "visible staff", "front stage", "visible"], "backstage": ["backstage actions", "back stage"], "support": ["support processes", "support process"]},
        "basis": {"observed": [], "assumed": [], "proposed": []},
    }
    if kind in choices:
        for canonical, aliases in choices[kind].items():
            if normalize(text) in {normalize(s) for s in [canonical]+aliases}:
                return canonical
        raise io.DataProblem("Use one of: " + ", ".join(choices[kind]) + ".")
    return text


def mapped_rows(frame, mapping, fields, comma=False, table="Table"):
    used = [c for c in mapping.values() if c is not None]
    if len(used) != len(set(used)):
        raise io.DataProblem(f"{table}: each source column can fill only one role. Check the column choices.")
    rows = []
    for index, source in enumerate(frame.to_dict("records"), 2):
        row = {}
        for field, (title, kind, _) in fields.items():
            column = mapping.get(field)
            try:
                row[field] = cast(source.get(column) if column else None, kind, comma)
            except io.DataProblem as exc:
                raise io.DataProblem(f"{table}, data row {index-1}, {title}: {exc}") from exc
        rows.append(row)
    return rows


def schema_rows(model, table, frame, mapping, comma=False):
    """Convert mapped workbook cells without supplying made-up numeric defaults."""
    specs = model.SCHEMA["properties"][table]["items"]["properties"]
    used = [v for v in mapping.values() if v]
    if len(used) != len(set(used)):
        raise io.DataProblem(f"{fmt.TABLE_NAMES[table]}: choose a different source column for each role.")
    rows = []
    for index, source in enumerate(frame.to_dict("records"), 1):
        row = {}
        for field, spec in specs.items():
            actual = spec.get("anyOf", [spec])[0]
            nullable = any(s.get("type") == "null" for s in spec.get("anyOf", []))
            value = source.get(mapping.get(field))
            try:
                if missing(value):
                    if nullable:
                        value = None
                    elif actual.get("type") == "string" and "minLength" not in actual and "pattern" not in actual and "format" not in actual:
                        value = ""
                    else:
                        raise io.DataProblem("Choose a column and supply a value.")
                elif actual.get("type") in {"number", "integer"}:
                    value = number(value, comma, field == "probability")
                    if actual["type"] == "integer":
                        if not value.is_integer():
                            raise io.DataProblem("Use a whole number.")
                        value = int(value)
                elif actual.get("format") == "date" and isinstance(value, (date, datetime)):
                    value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                elif "enum" in actual:
                    matched = suggest(actual["enum"], [str(value)])
                    if matched is None:
                        raise io.DataProblem("Use one of: " + ", ".join(actual["enum"]) + ".")
                    value = matched
                else:
                    value = str(value).strip()
                row[field] = value
            except io.DataProblem as exc:
                raise io.DataProblem(f"{fmt.TABLE_NAMES[table]}, data row {index}, {label(field)}: {exc}") from exc
        rows.append(row)
    return rows


def complete_project(model, brief, tables):
    d = {"schema_version": "1.0", "brief": brief}
    for name in model.TITLES:
        d[name] = deepcopy(tables.get(name, []))
        minimum = model.SCHEMA["properties"][name].get("minItems", 0)
        if len(d[name]) < minimum:
            raise io.DataProblem(f"{fmt.TABLE_NAMES[name]} needs at least {minimum} row(s). Choose its sheet and check the columns.")
    return model.validate(d)


def _workbook(tables, instructions):
    book = Workbook()
    book.remove(book.active)
    data = {"Read me": pd.DataFrame({"Instructions": instructions}), **tables}
    for title, frame in data.items():
        safe_title = re.sub(r"[\\/*?:\[\]]", " ", title)[:31]
        sheet = book.create_sheet(safe_title)
        for row in [list(frame.columns)] + frame.astype(object).where(pd.notna(frame), None).values.tolist():
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Preserve literal text without creating an Excel formula.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="343229")
            cell.font = Font(color="FFFFFF", bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for i, column in enumerate(frame.columns, 1):
            values = [str(column)] + [str(v) for v in frame[column].head(30) if v is not None]
            sheet.column_dimensions[get_column_letter(i)].width = min(60, max(18, max(map(len, values), default=18)+2))
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def simple_template():
    tables = {}
    for key, spec in fmt.QUICK.items():
        tables[spec["title"]] = pd.DataFrame(fmt.EXAMPLES[key], columns=list(spec["fields"])).rename(columns={k: v[0] for k, v in spec["fields"].items()})
    return _workbook(tables, ["FICTIONAL EXAMPLE — replace every example row with your own data before using results.", fmt.INTRO, fmt.HELP, "Keep the first row as column headings. Blank numeric cells mean unknown, not zero.", "Upload using Simple business tables. Confirm the sheet and column suggestions before importing."])


def csv_template(key):
    spec = fmt.QUICK[key]
    frame = pd.DataFrame(fmt.EXAMPLES[key], columns=list(spec["fields"])).rename(columns={k: v[0] for k, v in spec["fields"].items()})
    return io.csv_bytes(frame)


def project_workbook(model, data, results=None):
    tables = {"Case": pd.DataFrame({"Case brief": [data["brief"]]})}
    for name in model.TITLES:
        columns = list(model.SCHEMA["properties"][name]["items"]["properties"])
        tables[fmt.TABLE_NAMES[name]] = pd.DataFrame(data[name], columns=columns).rename(columns={c: label(c) for c in columns})
    for i, (name, frame) in enumerate((results or {}).items(), 1):
        if name not in model.TITLES:
            tables[f"Result {i} {name}"[:31]] = frame
    return _workbook(tables, ["Signal project workbook. Imported workbooks always need a new review; review signatures are not imported from Excel.", "If this is the fictional example, replace example rows and the case brief before using it for your business.", "Import using Complete project workbook. Sheets and columns are suggested automatically; confirm or correct them.", "Reference columns link related tables. Preserve references when renaming a record. Empty numeric cells remain unknown.", "Result sheets are exports only; the importer does not use them as model inputs."])
