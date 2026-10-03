from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from blueprintsignal import model, input_format as fmt, portable as io, spreadsheets as sheets
from blueprintsignal.ui import data_input

APP = Path(__file__).resolve().parents[1] / "app.py"
NAME = "blueprint"


def read_complete(raw):
    tables = sheets.load_tables([("project.xlsx", raw)])
    mapped = {}
    for key in model.TITLES:
        frame = tables[fmt.TABLE_NAMES[key]]
        columns = model.SCHEMA["properties"][key]["items"]["properties"]
        mapping = {field: sheets.suggest(list(frame), [sheets.label(field), field]) for field in columns}
        mapped[key] = sheets.schema_rows(model, key, frame, mapping)
    return sheets.complete_project(model, tables["Case"].iloc[0]["Case brief"], mapped)


def quick_tables(raw):
    tables = sheets.load_tables([("simple.xlsx", raw)])
    mapped = {}
    for key, spec in fmt.QUICK.items():
        frame = tables[spec["title"]]
        mapping = {field: sheets.suggest(list(frame), [title, field]+aliases) for field, (title, _, aliases) in spec["fields"].items()}
        mapped[key] = sheets.mapped_rows(frame, mapping, spec["fields"])
    return mapped


def assert_same_inputs(actual, expected):
    # Excel stores numeric cells at decimal precision; allow only rounding noise.
    if isinstance(expected, float):
        assert actual == pytest.approx(expected, rel=1e-14, abs=1e-14)
    elif isinstance(expected, dict):
        assert actual.keys() == expected.keys()
        for key in expected:
            assert_same_inputs(actual[key], expected[key])
    elif isinstance(expected, list):
        assert len(actual) == len(expected)
        for a, e in zip(actual, expected):
            assert_same_inputs(a, e)
    else:
        assert actual == expected


def test_full_excel_roundtrip_preserves_all_inputs_and_no_review():
    d = model.demo()["data"]
    assert_same_inputs(read_complete(sheets.project_workbook(model, d)), d)
    p = io.project(NAME, read_complete(sheets.project_workbook(model, d)), "Spreadsheet draft")
    assert not io.reviewed(p)


def test_blank_workbook_roundtrip_keeps_unknowns():
    d = model.starter("Our own case")
    assert read_complete(sheets.project_workbook(model, d)) == d


def test_simple_example_builds_valid_independent_case():
    mapped = quick_tables(sheets.simple_template())
    d = fmt.build("Our imported case", mapped, {}, "example.xlsx")
    assert model.validate(d) == d
    assert d["brief"] == "Our imported case"
    assert d["sources"][0]["id"] == "FILE1"
    assert all(r.get("source_id", "FILE1") != "N1" for k in model.TITLES for r in d[k])
    assert [x["name"] for x in d["stages"]] == ["Arrive", "Order"]
    assert len(d["actions"]) == 3 and d["actions"][0]["minutes"] is None
    assert not d["links"]


def test_csv_keeps_text_and_explicit_decimal_choice():
    raw = 'name;value;note\nNA;1,5;unknown\nB;;'.encode()
    frame = sheets.load_tables([("example.csv", raw)])["example"]
    fields = {"name": ("Name", "text", []), "value": ("Value", "optional_number", []), "note": ("Note", "optional_text", [])}
    mapping = {k: k for k in fields}
    rows = sheets.mapped_rows(frame, mapping, fields, comma=True)
    assert rows == [{"name": "NA", "value": 1.5, "note": "unknown"}, {"name": "B", "value": None, "note": ""}]
    with pytest.raises(io.DataProblem, match="comma decimals"):
        sheets.mapped_rows(frame, mapping, fields)


@pytest.mark.parametrize("value", ["forty", "NaN", "Infinity", True, "40"])
def test_invalid_probabilities_are_not_replaced_with_zero(value):
    with pytest.raises(io.DataProblem):
        sheets.cast(value, "optional_probability")


def test_percentage_and_column_aliases():
    assert sheets.cast("40%", "optional_probability") == .4
    assert sheets.suggest(["Other", "Customer ID"], ["customer_id"]) == "Customer ID"
    assert sheets.suggest(["Other"], ["missing"]) is None
    assert sheets.cast("Visible staff / technology", "lane") == "frontstage"
    assert sheets.cast("", "basis") == "assumed"


@pytest.mark.parametrize("raw", [b"a,a\n1,2", b"a,\n1,2", b"a,b\n1,2,3"])
def test_bad_headers_or_ragged_csv_are_rejected(raw):
    with pytest.raises(io.DataProblem):
        sheets.load_tables([("bad.csv", raw)])


def test_duplicate_mapping_rejected():
    frame = pd.DataFrame({"a": [1]})
    fields = {"x": ("X", "number", []), "y": ("Y", "number", [])}
    with pytest.raises(io.DataProblem, match="only one role"):
        sheets.mapped_rows(frame, {"x": "a", "y": "a"}, fields)


def test_uncached_formula_requires_saved_values():
    book = Workbook()
    book.active.append(["Amount"])
    book.active.append(["=1+1"])
    out = BytesIO()
    book.save(out)
    with pytest.raises(io.DataProblem, match="no saved result"):
        sheets.load_tables([("formula.xlsx", out.getvalue())])


def test_formula_like_notes_export_as_text_and_roundtrip():
    d = model.demo()["data"]
    d["sources"][0]["note"] = '=HYPERLINK("https://example.org","test")'
    raw = sheets.project_workbook(model, d)
    book = load_workbook(BytesIO(raw), data_only=False)
    assert all(cell.data_type != "f" for row in book["Sources"] for cell in row)
    assert_same_inputs(read_complete(raw), d)


def test_multiple_csv_files_and_example_csvs():
    files = [(spec["title"]+".csv", sheets.csv_template(key)) for key, spec in fmt.QUICK.items()]
    tables = sheets.load_tables(files)
    assert len(tables) == len(fmt.QUICK)
    assert all(len(frame) for frame in tables.values())


def test_complete_missing_required_sheet_rejected():
    with pytest.raises(io.DataProblem, match="needs at least"):
        sheets.complete_project(model, "A case", {})


def test_invalid_files_are_friendly():
    for name, raw in [("bad.xlsx", b"not a workbook"), ("bad.xls", b"anything"), ("bad.csv", b"\xff\xfe")]:
        with pytest.raises(io.DataProblem):
            sheets.load_tables([(name, raw)])


def test_start_choices_file_default_and_manual_navigation():
    a = AppTest.from_file(str(APP), default_timeout=30).run()
    a.button(key=NAME+":start:Excel or CSV").click().run()
    assert a.sidebar.radio[0].value == "1 · Add your data"
    assert a.radio(key=NAME+":input_mode").value == "Excel or CSV"
    assert not a.exception and not a.error
    a.radio(key=NAME+":input_mode").set_value("Enter manually").run()
    next(t for t in a.text_area if t.label == "Your business question").set_value("Our own manually entered case")
    next(b for b in a.button if b.label == "Start entering my data").click().run()
    assert not a.exception and not a.error
    assert a.sidebar.radio[0].value == "2 · Edit & review"
    p = a.session_state[NAME+":project"]
    assert p["data"]["brief"] == "Our own manually entered case"
    assert not io.reviewed(p)


def test_ai_start_choice_still_works():
    a = AppTest.from_file(str(APP), default_timeout=30).run()
    a.button(key=NAME+":start:Use your AI").click().run()
    assert a.radio(key=NAME+":input_mode").value == "Use your AI"
    assert a.text_area(key=NAME+":ai_json")
    assert not a.exception and not a.error


@pytest.mark.parametrize("complete", [False, True])
def test_upload_mapping_preview_and_commit_clear_review(monkeypatch, complete):
    uploaded = BytesIO(sheets.project_workbook(model, model.demo()["data"]) if complete else sheets.simple_template())
    uploaded.name = "our-data.xlsx"
    monkeypatch.setattr(data_input.st, "file_uploader", lambda *args, **kwargs: [uploaded])
    a = AppTest.from_file(str(APP), default_timeout=30).run()
    original = a.session_state[NAME+":project"]
    a.button(key=NAME+":start:Excel or CSV").click().run()
    assert not a.exception
    assert a.session_state[NAME+":project"] == original
    if not complete:
        next(t for t in a.text_area if t.label == "What question should these data help answer?").set_value("Test import of our business case")
        a.run()
    assert not a.exception and not a.error and not a.warning
    a.button(key=NAME+":use_spreadsheet").click().run()
    assert not a.exception and not a.error
    assert a.sidebar.radio[0].value == "2 · Edit & review"
    assert not io.reviewed(a.session_state[NAME+":project"])
    if complete:
        assert_same_inputs(a.session_state[NAME+":project"]["data"], original["data"])


def test_bad_upload_does_not_change_current_project(monkeypatch):
    uploaded = BytesIO(b"heading,heading\n1,2")
    uploaded.name = "broken.csv"
    monkeypatch.setattr(data_input.st, "file_uploader", lambda *args, **kwargs: [uploaded])
    a = AppTest.from_file(str(APP), default_timeout=30).run()
    original = a.session_state[NAME+":project"]
    a.button(key=NAME+":start:Excel or CSV").click().run()
    assert a.error and not a.exception
    assert a.session_state[NAME+":project"] == original


def test_file_limit_follows_the_50_mb_upload_cap():
    assert sheets.MAX_BYTES == 50 * 1024 * 1024
    note = "x" * 700
    raw = ("Stage,Service layer,What happens\n" + "".join(f"S,Customer actions,{i} {note}\n" for i in range(9_000))).encode()
    assert 5_000_000 < len(raw) < sheets.MAX_BYTES  # above the old 5 MB limit, well inside the new one
    assert len(sheets.load_tables([("large.csv", raw)])["large"]) == 9_000
    with pytest.raises(io.DataProblem, match="no more than 1 MB"):
        sheets.load_tables([("large.csv", raw)], max_bytes=1024 * 1024)


def test_row_guard_is_a_method_limit_with_a_clear_message():
    raw = ("Stage,What happens\n" + "".join(f"S,Step {i}\n" for i in range(sheets.MAX_ROWS + 5))).encode()
    with pytest.raises(io.DataProblem, match="one service"):
        sheets.load_tables([("long.csv", raw)])
